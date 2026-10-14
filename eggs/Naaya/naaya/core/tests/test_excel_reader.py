import datetime
import unittest
from io import BytesIO

import openpyxl
import xlwt

from naaya.core.zope2util import read_excel_rows, new_xlsx_workbook

ROWS = [
    ['Title', 'Count', 'Note'],
    ['First', 3, 'caf\xe9'],
    ['Second', 2.5, ''],
]
EXPECTED = [
    ['Title', 'Count', 'Note'],
    ['First', 3.0, 'caf\xe9'],
    ['Second', 2.5, ''],
]


def make_xlsx(rows):
    output, wb = new_xlsx_workbook()
    ws = wb.add_worksheet('Data')
    wb.add_worksheet('Other').write(0, 0, 'ignored')
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            ws.write(r, c, value)
    wb.close()
    return output.getvalue()


def make_xls(rows):
    wb = xlwt.Workbook(encoding='utf-8')
    ws = wb.add_sheet('Data')
    wb.add_sheet('Other').write(0, 0, 'ignored')
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            ws.write(r, c, value)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


class ReadExcelRowsTest(unittest.TestCase):

    def test_xlsx(self):
        data = make_xlsx(ROWS)
        self.assertEqual(data[:2], b'PK')
        self.assertEqual(read_excel_rows(data), EXPECTED)

    def test_xls(self):
        data = make_xls(ROWS)
        self.assertNotEqual(data[:2], b'PK')
        self.assertEqual(read_excel_rows(data), EXPECTED)

    def test_same_rows_for_both_formats(self):
        self.assertEqual(read_excel_rows(make_xlsx(ROWS)),
                         read_excel_rows(make_xls(ROWS)))

    def test_file_object(self):
        self.assertEqual(read_excel_rows(BytesIO(make_xlsx(ROWS))), EXPECTED)

    def test_short_rows_are_padded(self):
        rows = [['A', 'B', 'C'], ['x']]
        expected = [['A', 'B', 'C'], ['x', '', '']]
        self.assertEqual(read_excel_rows(make_xlsx(rows)), expected)
        self.assertEqual(read_excel_rows(make_xls(rows)), expected)

    def test_trailing_empty_rows_are_dropped(self):
        rows = [['A'], ['x'], [''], ['']]
        self.assertEqual(read_excel_rows(make_xlsx(rows)), [['A'], ['x']])
        self.assertEqual(read_excel_rows(make_xls(rows)), [['A'], ['x']])

    def test_dates_are_excel_serials(self):
        output, wb = new_xlsx_workbook()
        ws = wb.add_worksheet()
        ws.write_datetime(0, 0, datetime.datetime(2020, 1, 31),
                          wb.add_format({'num_format': 'dd/mm/yyyy'}))
        wb.close()
        self.assertEqual(read_excel_rows(output.getvalue()), [[43861.0]])

    def test_invalid_file(self):
        self.assertRaises(ValueError, read_excel_rows, b'not an excel file')
        self.assertRaises(ValueError, read_excel_rows, b'PKnot a zip')
        self.assertRaises(ValueError, read_excel_rows, b'')

    def test_writer_output_opens_in_openpyxl(self):
        wb = openpyxl.load_workbook(BytesIO(make_xlsx(ROWS)))
        self.assertEqual(wb.sheetnames, ['Data', 'Other'])


def test_suite():
    return unittest.TestSuite([
        unittest.TestLoader().loadTestsFromTestCase(ReadExcelRowsTest)])
