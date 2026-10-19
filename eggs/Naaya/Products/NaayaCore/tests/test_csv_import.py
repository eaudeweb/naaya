from unittest import TestSuite, TestLoader
from unittest.mock import Mock, patch
from io import StringIO, BytesIO

from zope import component, interface
from zope.component.globalregistry import getGlobalSiteManager

from Products.Naaya.tests.NaayaTestCase import NaayaTestCase
from Products.Naaya.tests.NaayaFunctionalTestCase import NaayaFunctionalTestCase
from Products.NaayaCore.SchemaTool.widgets.geo import Geo
from Products.NaayaCore.interfaces import ICSVImportExtraColumns
from naaya.content.url.interfaces import INyURL
from Products.NaayaCore.EmailTool import EmailTool
from Products.NaayaCore.EmailTool.EmailPageTemplate import EmailPageTemplateFile
from Products.Naaya.NyFolder import addNyFolder

from Products.NaayaCore.GeoMapTool.tests.test_kml_parser import load_file
from naaya.core.zope2util import XLSX_CONTENT_TYPE
from naaya.core.tests.test_excel_reader import make_xlsx, make_xls

csv_data = ('Title,Description,Automatically redirect to the given URL,URL\n'
    'My URL 1,The best URL,no,http://example.com\n'
    'Eau de Web,The Naaya company,yes,http://www.eaudeweb.ro\n')

def check_uploaded(test):
    test.assertEqual(len(test.portal.imported.objectIds()), 2)
    test.assertEqual(set(map(lambda url: url.title, test.portal.imported.objectValues())),
        set(['My URL 1', 'Eau de Web']))
    test.assertEqual(set(map(lambda url: url.locator, test.portal.imported.objectValues())),
        set(['http://example.com', 'http://www.eaudeweb.ro']))
    test.assertEqual(test.portal.imported._getOb('my-url-1').redirect, False)
    test.assertEqual(test.portal.imported._getOb('eau-de-web').redirect, True)

def change_encoding(string, current, new):
    if isinstance(string, str):
        string = string.encode(current)
    return string.decode(current).encode(new)

def do_import_object(context, meta_type, csv_data, target):
    """ Does the csv import in the given target with the given data.
        Target must exist.
    """
    target = context.portal.restrictedTraverse(target, None)
    target.csv_import.do_import(meta_type=meta_type, file_type='CSV',
                                data=StringIO(csv_data))

class NyCSVImportTest(NaayaTestCase):
    """ TestCase for Naaya CSV import """

    def afterSetUp(self):
        addNyFolder(self.portal, 'imported', contributor='contributor', submitted=1)

    def beforeTearDown(self):
        self.portal.manage_delObjects(['imported'])

    def test_generate_csv_template(self):
        columns = self.portal.csv_import.template('Naaya URL',
                                                  'CSV').strip().split(',')
        self.assertEqual(len(columns), 14)
        self.assertTrue('Title' in columns)
        self.assertTrue('Description' in columns)
        self.assertTrue('Automatically redirect to the given URL' in columns)
        self.assertTrue('URL' in columns)

    def test_import_ok(self):
        do_import_object(self, 'Naaya URL', csv_data, 'imported')
        check_uploaded(self)

    def test_import_document(self):
        data = ('Title,Description,Geographical coverage,Keywords,Sort order,'
                'Release date,Open for comments,Body (HTML)\n'
            'My doc,,,,,,,This is a test document.\n')
        do_import_object(self, 'Naaya Document', data, 'imported')
        self.assertTrue('my-doc' in self.portal.imported.objectIds())
        my_doc = self.portal.imported._getOb('my-doc')
        self.assertEqual(my_doc.meta_type, 'Naaya Document')
        self.assertEqual(my_doc.title, 'My doc')
        self.assertEqual(my_doc.body, 'This is a test document.')
        self.assertEqual(my_doc.sortorder, 100)
        self.assertTrue(my_doc.approved)

    def test_import_unicode_document(self):
        data = ('Title,Description,Geographical coverage,Keywords,Sort order,'
                'Release date,Open for comments,Body (HTML)\n'
            'Forschungsinstitut f\xfcr Freizeit und Tourismus (FIF),,,,,,,This is a test document.\n')
        do_import_object(self, 'Naaya Document', data, 'imported')
        self.assertTrue('forschungsinstitut-fur-freizeit-und-tourismus-fif' in self.portal.imported.objectIds())
        my_doc = self.portal.imported._getOb('forschungsinstitut-fur-freizeit-und-tourismus-fif')
        self.assertEqual(my_doc.meta_type, 'Naaya Document')
        self.assertEqual(my_doc.title, 'Forschungsinstitut f\xfcr Freizeit und Tourismus (FIF)')
        self.assertEqual(my_doc.body, 'This is a test document.')
        self.assertEqual(my_doc.sortorder, 100)
        self.assertTrue(my_doc.approved)

    def test_import_bad_unicode_document(self):
        """ In Python 3 all strings are unicode, so 'bad encoding' at the
        string level doesn't apply.  Instead, verify that importing a
        document with non-ASCII characters in the title works. """
        data = ('Title,Description,Geographical coverage,Keywords,Sort order,'
                'Release date,Open for comments,Body (HTML)\n'
                'Forschungsinstitut f\xfcr Freizeit und Tourismus (FIF),'
                ',,,,,,This is a test document.\n')
        do_import_object(self, 'Naaya Document', data, 'imported')
        self.assertTrue(
            'forschungsinstitut-fur-freizeit-und-tourismus-fif'
            in self.portal.imported.objectIds())

    def test_import_bad_data(self):
        def do_import(row=''):
            data = csv_data + row
            do_import_object(self, 'Naaya URL', data, 'imported')
        try:
            do_import('T,,,http://example.com\n')
        except:
            self.fail('Should not raise exception')

        self.assertEqual(len(self.portal.imported.objectIds()), 3)
        self.assertEqual(self.portal.imported._getOb('t').title, 'T')

        self.assertRaises(ValueError, do_import, ',D,yes,http://example.com\n')
        self.assertRaises(ValueError, do_import, 'T,D,asdf,http://example.com\n')

    def test_import_bad_metatype(self):
        def do_import():
            do_import_object(self, 'Nonexistent Metatype', csv_data, 'imported')
        self.assertRaises(ValueError, do_import)

    def test_extra_csv_columns(self):
        csv_with_extra = ('Title,something,something_else\n'
                          'TY,asdf,qwer\n')
        extra_data = []

        @interface.implementer(ICSVImportExtraColumns)
        @component.adapter(INyURL)
        class UrlAdapter(object):
            def __init__(self, ob):
                self.ob = ob
            def handle_columns(self, extra_properties):
                extra_data.append(extra_properties)

        reg = getGlobalSiteManager()
        reg.registerAdapter(UrlAdapter)

        do_import_object(self, 'Naaya URL', csv_with_extra, 'imported')

        self.assertEqual(self.portal.imported['ty'].title, 'TY')
        self.assertEqual(extra_data, [{'something': 'asdf',
                                       'something_else': 'qwer'}])

        reg.unregisterAdapter(UrlAdapter)

    def test_import_mails(self):
        diverted_mail = EmailTool.divert_mail()

        #Enable instant notifications
        notification_tool = self.portal.getNotificationTool()
        notification_tool.config['enable_instant'] = True
        self.portal.getNotificationTool().add_account_subscription(
            'contributor', '', 'instant', 'en')
        self.portal.imported.maintainer_email = 'someone@somehost'
        do_import_object(self, 'Naaya URL', csv_data, 'imported')

        self.assertEqual(len(diverted_mail), 3)

        expected_subject = u'CSV Import - imported'
        expected_body = (u'This is automatically generated message'
                          ' to inform you that the following 2 items'
                          ' have been uploaded in imported'
                          ' (http://nohost/portal/imported):\n'
                          ' - My URL 1\n - Eau de Web\n\n'
                          'Uploaded by Anonymous User on')
        expected_recipients = set(['site.admin@example.com',  # administrator
                                   'someone@somehost',  # folder_maintainer
                                   'contrib@example.com'])  # subscriber
        expected_sender = 'from.zope@example.com'

        actual_recipients = set(m[1][0] for m in diverted_mail)
        self.assertEqual(actual_recipients, expected_recipients)
        # Check body and sender on any mail (they're the same for CSV import)
        csv_mails = [m for m in diverted_mail if m[4] == expected_subject]
        self.assertTrue(len(csv_mails) > 0)
        self.assertTrue(expected_body in csv_mails[0][0])
        self.assertEqual(expected_sender, csv_mails[0][3])
        EmailTool.divert_mail(False)

    def test_import_default_values(self):
        data = 'Title,Sort order\nMy doc,\n'
        do_import_object(self, 'Naaya Document', data, 'imported')
        self.assertTrue('my-doc' in self.portal.imported.objectIds())
        my_doc = self.portal.imported._getOb('my-doc')
        self.assertEqual(my_doc.meta_type, 'Naaya Document')
        self.assertEqual(my_doc.title, 'My doc')
        self.assertEqual(my_doc.sortorder, 100)
        self.assertTrue(my_doc.approved)

excel_rows = [
    ['Title', 'Description', 'Automatically redirect to the given URL', 'URL'],
    ['My URL 1', 'The best URL', 'no', 'http://example.com'],
    ['Eau de Web', 'The Naaya company', 'yes', 'http://www.eaudeweb.ro'],
]


class NyExcelImportTest(NaayaTestCase):
    """ Excel import, with .xlsx and .xls files """

    def afterSetUp(self):
        addNyFolder(self.portal, 'imported', contributor='contributor',
                    submitted=1)

    def beforeTearDown(self):
        self.logout()
        self.portal.manage_delObjects(['imported'])

    def do_import(self, data, **kwargs):
        return self.portal.imported.csv_import.do_import(
            meta_type='Naaya URL', file_type='Excel', data=BytesIO(data),
            **kwargs)

    def test_import_xlsx(self):
        self.do_import(make_xlsx(excel_rows))
        check_uploaded(self)

    def test_import_xls(self):
        self.do_import(make_xls(excel_rows))
        check_uploaded(self)

    def test_import_unicode(self):
        rows = [['Title', 'URL'], ['Forschungsinstitut f\xfcr Freizeit',
                                   'http://example.com']]
        for make in (make_xlsx, make_xls):
            self.do_import(make(rows), overwrite_existing=True)
            ob = self.portal.imported._getOb(
                'forschungsinstitut-fur-freizeit')
            self.assertEqual(ob.title, 'Forschungsinstitut f\xfcr Freizeit')
            self.portal.imported.manage_delObjects([ob.getId()])

    def test_import_truncated_cells(self):
        from Products.NaayaCore.managers.import_export import CSVImportTool
        mark = ('[TRUNCATED: the original text had 500 characters; '
                'Excel cells hold at most 100]')
        rows = [['ID', 'Title', 'URL', 'Description'],
                ['u1', 'T1', 'http://example.com', 'text ' + mark],
                ['u2', 'T2', 'http://example.com', 'full text']]
        self.login()
        with patch.object(type(self.portal.aq_base), 'setSessionErrorsTrans') as errors, \
                patch.object(CSVImportTool, 'index_html'):
            self.do_import(make_xlsx(rows), REQUEST=self.portal.REQUEST)
        [(message, mapping)] = errors.call_args[0][0]
        self.assertIn('imported with text truncated', message)
        self.assertEqual(mapping, {'count': 1, 'rows': 'u1 (Description)'})
        self.assertEqual(self.portal.imported.u1.description, 'text ' + mark)
        self.assertEqual(sorted(self.portal.imported.objectIds()), ['u1', 'u2'])

    def test_invalid_file(self):
        self.assertRaises(ValueError, self.do_import, b'not an excel file')
        self.assertRaises(ValueError, self.do_import, b'PKnot a zip')
        self.assertEqual(self.portal.imported.objectIds(), [])

    def test_generate_excel_template(self):
        import openpyxl
        self.login('site_admin')
        request = Mock()
        ret = self.portal.csv_import.template(
            'Naaya URL', 'Excel', as_attachment=True, REQUEST=request)
        ws = openpyxl.load_workbook(BytesIO(ret)).worksheets[0]
        header = [c.value for c in ws[1]]
        self.assertEqual(header[0], 'ID')
        self.assertTrue('Title' in header)
        self.assertTrue('URL' in header)
        self.assertTrue(
            ws['A1'].font.b, 'template header should be bold')
        response = request.RESPONSE
        response.setHeader.assert_any_call(
            'Content-Type', XLSX_CONTENT_TYPE)
        disposition = [c.args[1] for c in response.setHeader.call_args_list
                       if c.args[0] == 'Content-Disposition'][0]
        self.assertTrue(disposition.endswith('.xlsx'), disposition)


class CSVImportFunctionalTests(NaayaFunctionalTestCase):

    def test_excel_upload(self):
        for make, name in ((make_xlsx, 'bulk.xlsx'), (make_xls, 'bulk.xls')):
            self.browser_do_login('admin', '')
            self.browser.go('http://localhost/portal/info/csv_import')
            form = self.browser.get_form('csv_import')
            form['meta_type'] = ['Naaya URL']
            form['file_type'] = ['Excel']
            form.find_control('data').add_file(
                BytesIO(make(excel_rows)), 'application/octet-stream', name)
            self.browser.clicked(form, form.find_control('do_import:method'))
            self.browser.submit()
            urls = self.portal.info.objectValues('Naaya URL')
            self.assertEqual(set(u.title for u in urls),
                             {'My URL 1', 'Eau de Web'})
            self.portal.info.manage_delObjects([u.getId() for u in urls])
            import transaction
            transaction.commit()


    def test_bad_encoding(self):
        # Produce a CSV with latin-1 encoded data that is NOT valid UTF-8.
        header = 'Title,Description,Geographical coverage,Keywords,Sort order,Release date,Open for comments,Body (HTML)\n'
        row = 'Forschungsinstitut f\xfcr Freizeit und Tourismus (FIF),,,,,,,This is a test document.\n'
        data = BytesIO(header.encode('ascii') + row.encode('latin-1'))
        before = self.portal.info.objectIds()
        self.browser_do_login('admin', '')
        self.browser.go('http://localhost/portal/info/csv_import')
        form = self.browser.get_form('csv_import')
        form['meta_type'] = ['Naaya Document']
        form.find_control('data').add_file(data, 'text/csv', 'HTML Document bulk upload')
        self.browser.clicked(form, form.find_control('do_import:method'))
        self.browser.submit()
        self.assertTrue('CSV file is not utf-8 encoded' in self.browser.get_html())
        self.assertEqual(before, self.portal.info.objectIds())

class GeopointImportTest(NaayaTestCase):
    def afterSetUp(self):
        addNyFolder(self.portal, 'imported', contributor='contributor', submitted=1)
        schema = self.portal.portal_schemas.NyDocument
        schema.addWidget('test_geo_loc', label="Geo Loc", widget_type='Geo', data_type='geo')
        schema.addWidget('test_geo_type', label="Geo Type", widget_type='GeoType', data_type='str')

        picture_data = load_file('data/symbol.png')

        self.portal.portal_map.addSymbol('sym1', 'Test symbol one', '', '', None, picture_data, '')
        self.portal.portal_map.addSymbol('sym2', 'Test symbol two', '', '', None, picture_data, '')

    def beforeTearDown(self):
        self.portal.portal_map.deleteSymbol(['sym1', 'sym2'])
        self.portal.portal_schemas.NyDocument.manage_delObjects('test_geo_loc-property')
        self.portal.portal_schemas.NyDocument.manage_delObjects('test_geo_type-property')
        self.portal.manage_delObjects(['imported'])

    def test_template(self):
        columns = self.portal.csv_import.template('Naaya Document',
                                                  'CSV').strip().split(',')
        self.assertTrue('Geo Loc - lat' in columns)
        self.assertTrue('Geo Loc - lon' in columns)
        self.assertTrue('Geo Loc - address' in columns)
        self.assertTrue('Geo Type' in columns)

    @patch('Products.NaayaCore.GeoMapTool.managers.geocoding.location_geocode',
           return_value=('44.434295', '26.102965'))
    def test_import(self, mock_geocode):
        geo_csv_data = (
            "Title,Geo Loc - lat,Geo Loc - lon,Geo Loc - address,Geo Type\n"
            "doc_one,,,,\n"
            "doc_two,13.45,22.60,,Test symbol one\n"
            "doc_three,8,9,somewhere else,Test symbol two\n"
            "doc_four,,,Bucharest,Test symbol two\n"
        )
        try:
            do_import_object(self, 'Naaya Document', geo_csv_data, 'imported')
        except:
            raise
            self.fail('Should not raise exception')
        self.assertEqual(len(self.portal.imported.objectIds()), 4)

        doc_one = self.portal.imported._getOb('doc_one')
        self.assertEqual(doc_one.title, 'doc_one')
        self.assertFalse(hasattr(doc_one, 'test_geo_loc'))
        self.assertFalse(hasattr(doc_one, 'test_geo_type'))

        doc_two = self.portal.imported._getOb('doc_two')
        self.assertEqual(doc_two.test_geo_loc, Geo('13.45', '22.60'))
        self.assertEqual(doc_two.test_geo_type, 'sym1')

        doc_three = self.portal.imported._getOb('doc_three')
        self.assertEqual(doc_three.test_geo_loc, Geo('8', '9', 'somewhere else'))
        self.assertEqual(doc_three.test_geo_type, 'sym2')

        doc_four = self.portal.imported._getOb('doc_four')
        correct = Geo('44.434295', '26.102965', 'Bucharest')
        self.assertTrue(-1 < (doc_four.test_geo_loc.lat - correct.lat)*100 < 1)
        self.assertTrue(-1 < (doc_four.test_geo_loc.lon - correct.lon)*100 < 1)
        self.assertEqual(doc_four.test_geo_loc.address, correct.address)
        self.assertEqual(doc_four.test_geo_type, 'sym2')

class SecurityTestCase(NaayaFunctionalTestCase):
    def afterSetUp(self):
        pass

    def beforeTearDown(self):
        pass

    def test_anonymous(self):
        self.browser.go('http://localhost/portal/info/csv_import')
        self.assertTrue(self.browser.get_url().startswith(
            'http://localhost/portal/login_html'))

        self.browser.go('http://localhost/portal/csv_export/export_json')
        self.assertFalse(self.browser.get_html().startswith('{'))
        self.assertTrue(self.browser.get_url().startswith(
            'http://localhost/portal/login_html'))

    def test_admin(self):
        self.browser_do_login('admin', '')

        self.browser.go('http://localhost/portal/info/csv_import')
        self.assertEqual(self.browser.get_url(),
            'http://localhost/portal/info/csv_import')

        self.browser.go('http://localhost/portal/csv_export/export_json')
        self.assertTrue(self.browser.get_html().startswith('{'))

        self.browser_do_logout()
