from io import BytesIO

import openpyxl

from Products.Naaya.tests.NaayaFunctionalTestCase import NaayaFunctionalTestCase
from Products.Naaya.NyFolder import addNyFolder
from naaya.core.zope2util import XLSX_CONTENT_TYPE


class ExcelDownloadsTestCase(NaayaFunctionalTestCase):
    """ Excel downloads of the notification tool, mail list and map """

    def check_download(self, url):
        self.browser_do_login('admin', '')
        self.browser.go(url)
        response = self.browser.result
        self.assertEqual(response.headers['content-type'], XLSX_CONTENT_TYPE)
        disposition = response.headers['content-disposition']
        self.assertTrue(disposition.strip('"').endswith('.xlsx'),
                        disposition)
        return openpyxl.load_workbook(BytesIO(response.content)).worksheets[0]

    def test_subscriptions(self):
        notif = self.portal.getNotificationTool()
        notif.add_account_subscription('contributor', '', 'instant', 'en')
        import transaction
        transaction.commit()
        ws = self.check_download(
            'http://localhost/portal/portal_notification/download'
            '?file_type=Excel')
        self.assertEqual([c.value for c in ws[1]],
                         ['User', 'Location', 'Notification type',
                          'Language'])
        self.assertEqual([c.value for c in ws[2]],
                         ['Contributor Test (contrib@example.com)',
                          'entire portal', 'instant', 'en'])
        self.assertTrue(ws['A1'].font.b)
        notif.remove_account_subscription('contributor', '', 'instant', 'en')
        transaction.commit()

    def test_bulk_mail_list(self):
        ws = self.check_download(
            'http://localhost/portal/admin_bulk_mail_list_export'
            '?headers=Subject,Recipients&keys=subject,recipients')
        self.assertEqual([c.value for c in ws[1]], ['Subject', 'Recipients'])

    def test_map_export(self):
        addNyFolder(self.portal, 'mapfolder', contributor='contributor',
                    submitted=1)
        import transaction
        transaction.commit()
        try:
            ws = self.check_download(
                'http://localhost/portal/portal_map/export_csv'
                '?meta_type=Naaya%20Folder&file_type=excel')
            self.assertEqual(ws['A1'].value, 'ID')
        finally:
            self.portal.manage_delObjects(['mapfolder'])
            transaction.commit()
