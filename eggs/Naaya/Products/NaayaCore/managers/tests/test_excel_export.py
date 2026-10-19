import datetime
import unittest
from io import BytesIO

import openpyxl

from Products.Naaya.tests.NaayaFunctionalTestCase import NaayaFunctionalTestCase
from naaya.core.zope2util import XLSX_CONTENT_TYPE


class ExcelExportTestCase(NaayaFunctionalTestCase):
    """ Excel export of a folder, produced as .xlsx """

    def afterSetUp(self):
        from Products.Naaya.NyFolder import addNyFolder
        from naaya.content.news.news_item import addNyNews
        addNyFolder(self.portal, 'myfolder', contributor='contributor', submitted=1)
        addNyNews(self.portal.myfolder, id='mynews', submitted=1, contributor='contributor', title='The News')
        import transaction; transaction.commit()

    def beforeTearDown(self):
        self.portal.manage_delObjects(['myfolder'])
        import transaction; transaction.commit()

    def test_excel_export(self):
        self.browser_do_login('admin', '')
        self.browser.go('http://localhost/portal/myfolder/csv_export/export?meta_type=Naaya%20News&file_type=Excel&as_attachment=True')
        response = self.browser.result
        self.assertEqual(response.headers['content-type'], XLSX_CONTENT_TYPE)
        self.assertTrue(response.headers['content-disposition']
                        .strip('"').endswith('.xlsx'))
        ws = openpyxl.load_workbook(BytesIO(response.content)).worksheets[0]
        header = [c.value for c in ws[1]]
        self.assertEqual(header[:2], ['ID', 'Title'])
        self.assertTrue(ws['A1'].font.b)
        row = [c.value for c in ws[2]]
        self.assertEqual(row[0], 'mynews')
        self.assertEqual(row[1], 'The News')

    def test_excel_export_truncated_cell(self):
        from unittest import mock
        from Products.NaayaCore.managers.import_export import ExportTool
        exporter = self.portal.myfolder.csv_export
        with mock.patch.object(ExportTool, "CELL_CHAR_LIMIT", 100):
            self.portal.myfolder.mynews.description = '<p>%s</p>' % ('x' * 300)
            data = exporter.export("Naaya News", "Excel")
        wb = openpyxl.load_workbook(BytesIO(data))
        ws = wb.worksheets[0]
        header = [c.value for c in ws[1]]
        cell = ws.cell(2, header.index('Description') + 1).value
        self.assertLessEqual(len(cell), 100)
        self.assertTrue(cell.startswith('<p>xxx'))
        self.assertTrue(cell.endswith(
            '<p>[TRUNCATED: the original text had 307 characters; '
            'Excel cells hold at most 100]</p>'))
        notes = wb['Truncated cells']
        self.assertEqual([c.value for c in notes[4]],
                         ['mynews', 'The News', 'Description', 307])


class GenerateExcelTestCase(unittest.TestCase):
    """ generate_excel writes a bold header, the rows and dd/mm/yyyy dates """

    def test_layout(self):
        from Products.NaayaCore.managers.import_export import generate_excel
        data = generate_excel(['ID', 'Title', 'Date'],
                              [['a', '=not a formula', '31/01/2020'],
                               ['b', 'caf\xe9', '']])
        wb = openpyxl.load_workbook(BytesIO(data))
        self.assertEqual(wb.sheetnames, ['Sheet 1'])
        ws = wb.worksheets[0]
        self.assertEqual([c.value for c in ws[1]], ['ID', 'Title', 'Date'])
        self.assertTrue(all(c.font.b for c in ws[1]))
        self.assertEqual(ws['B2'].value, '=not a formula')
        self.assertEqual(ws['B2'].data_type, 's')
        self.assertEqual(ws['C2'].value, datetime.datetime(2020, 1, 31))
        self.assertEqual(ws['C2'].number_format, 'dd/mm/yyyy')
        self.assertEqual(ws['B3'].value, 'caf\xe9')
        self.assertFalse(ws['A2'].font.b)

    def test_template_without_rows(self):
        from Products.NaayaCore.managers.import_export import generate_excel
        data = generate_excel(['ID', 'Title'], [[]])
        ws = openpyxl.load_workbook(BytesIO(data)).worksheets[0]
        self.assertEqual([c.value for c in ws[1]], ['ID', 'Title'])


class TruncateCellTestCase(unittest.TestCase):

    def test_plain_text(self):
        from Products.NaayaCore.managers.import_export import truncate_cell
        cell = truncate_cell('a' * 200, 120)
        self.assertEqual(len(cell), 120)
        self.assertTrue(cell.endswith(' [TRUNCATED: the original text had 200 '
                                      'characters; Excel cells hold at most 120]'))

    def test_html_stays_well_formed(self):
        import lxml.html
        from Products.NaayaCore.managers.import_export import truncate_cell
        text = '<table><tr><td>' + 'cell</td><td>' * 50 + '</td></tr></table>'
        for limit in range(110, 140):
            cell = truncate_cell(text, limit)
            self.assertLessEqual(len(cell), limit)
            body = cell[:cell.rindex('\n<p>[TRUNCATED')]
            self.assertTrue(body.startswith('<table>'))
            self.assertTrue(body.endswith('</table>'))
            self.assertNotIn('<td', body[-6:])
            lxml.html.fragment_fromstring(cell, create_parent='div')
