import csv
from io import BytesIO, StringIO
from unittest.mock import Mock, patch

import openpyxl
from PIL import Image
from Testing import ZopeTestCase

from Products.Naaya.tests.NaayaTestCase import NaayaTestCase
from Products.NaayaSurvey.MegaSurvey import manage_addMegaSurvey
from naaya.core.zope2util import XLSX_CONTENT_TYPE

ZopeTestCase.installProduct('NaayaWidgets')
ZopeTestCase.installProduct('NaayaSurvey')


def png_bytes():
    output = BytesIO()
    Image.new('RGB', (40, 40), (200, 30, 30)).save(output, 'PNG')
    return output.getvalue()


class SurveyExcelExportTestCase(NaayaTestCase):
    """ Answers and report statistics are exported as .xlsx """

    def afterSetUp(self):
        self.login()
        self.portal.manage_install_pluggableitem('Naaya Mega Survey')
        survey_id = manage_addMegaSurvey(self.portal, title='Excel survey')
        self.survey = survey = self.portal._getOb(survey_id)
        self.radio = survey.addWidget(
            title='Radio q', meta_type='Naaya Radio Widget',
            choices=['Yes', 'No'])
        self.matrix = survey.addWidget(
            title='Matrix q', meta_type='Naaya Radio Matrix Widget',
            rows=['Row A', 'Row B'], choices=['Low', 'High'])
        self.combo = survey.addWidget(
            title='Combo q', meta_type='Naaya Combobox Matrix Widget',
            rows=['Row C'], choices=['Col 1', 'Col 2'])
        survey._getOb(self.combo).set_localproperty(
            'values', 'lines', 'en', ['Val 0', 'Val 1', 'Val 2'])
        self.text = survey.addWidget(
            title='Text q', meta_type='Naaya String Widget')
        self.date = survey.addWidget(
            title='Date q', meta_type='Naaya Date Widget')
        for radio, text in ((0, 'first answer'), (1, 'caf\xe9 answer')):
            form = {
                '%s' % self.radio: radio,
                '%s_0' % self.matrix: 0,
                '%s_1' % self.matrix: 1,
                '%s_0_0' % self.combo: 2,
                '%s_0_1' % self.combo: 1,
                '%s' % self.text: text,
                '%s' % self.date: '31/01/2020',
            }
            survey.addSurveyAnswer(notify_respondent=False, **form)

    def beforeTearDown(self):
        self.logout()

    def check_response(self, request, extension='.xlsx'):
        response = request.RESPONSE
        response.setHeader.assert_any_call('Content-Type', XLSX_CONTENT_TYPE)
        calls = [c.args for c in response.setHeader.call_args_list]
        disposition = [c[1] for c in calls
                       if c[0] == 'Content-Disposition'][0]
        self.assertTrue(disposition.endswith(extension), disposition)

    def test_download_answers(self):
        response = Mock()
        data = self.survey.download(
            REQUEST={'file_type': 'Excel'}, RESPONSE=response)
        response.setHeader.assert_any_call('Content-Type', XLSX_CONTENT_TYPE)
        disposition = response.setHeader.call_args_list[1].args[1]
        self.assertEqual(disposition,
                         'attachment; filename=%s.xlsx' % self.survey.id)
        ws = openpyxl.load_workbook(BytesIO(data)).worksheets[0]
        header = [c.value for c in ws[1]]
        self.assertEqual(header[0], 'Respondent')
        self.assertTrue('Radio q' in header)
        self.assertTrue(all(c.font.b for c in ws[1]))
        texts = {c.value for row in ws.iter_rows(min_row=2) for c in row}
        self.assertTrue({'first answer', 'caf\xe9 answer'} <= texts)

    def test_download_csv_unchanged(self):
        response = Mock()
        data = self.survey.download(
            REQUEST={'file_type': 'CSV'}, RESPONSE=response)
        response.setHeader.assert_any_call('Content-Type', 'text/csv')
        rows = list(csv.reader(StringIO(data)))
        self.assertEqual(rows[0][:4], ['Respondent', 'Radio q', 'Matrix q',
                                       'Row A'])
        self.assertEqual(len(rows), 3)
        self.assertEqual(sorted(row[rows[0].index('Text q')]
                                for row in rows[1:]),
                         ['caf\xe9 answer', 'first answer'])
        self.assertEqual({row[1] for row in rows[1:]}, {'Yes', 'No'})

    @patch('Products.NaayaSurvey.statistics.MultipleChoiceGoogleBarChart'
           'Statistic.urlopen')
    @patch('Products.NaayaSurvey.statistics.MultipleChoicePieChart'
           'Statistic.urlopen')
    def test_export_report(self, pie_urlopen, bar_urlopen):
        for mock in (pie_urlopen, bar_urlopen):
            mock.return_value.read.return_value = png_bytes()
        report_id = self.survey.generateFullReport(title='Full report')
        request = Mock()
        data = self.survey.questionnaire_export(report_id, request)
        self.check_response(request)
        wb = openpyxl.load_workbook(BytesIO(data))
        self.assertEqual(wb.sheetnames, ['Report'])
        ws = wb['Report']
        rows = [[c.value for c in row] for row in ws.iter_rows()]
        flat = [v for row in rows for v in row if v is not None]
        for text in ('Radio q', 'Matrix q', 'Combo q', 'Text q', 'Date q',
                     'Choice', 'Count', 'Percent', 'Answered', 'Not answered',
                     'Total', 'Yes', 'No', 'Row A', 'Low', 'High', 'Row C',
                     'Val 1', 'first answer', 'caf\xe9 answer'):
            self.assertTrue(text in flat, text)
        self.assertTrue('2 (100.00%)' in flat)
        self.assertTrue('50.00%' in flat)
        self.assertTrue('100%' in flat)
        link_cell = [c for row in ws.iter_rows() for c in row
                     if c.hyperlink is not None]
        self.assertEqual(len(link_cell), 2)
        header_cell = [c for row in ws.iter_rows() for c in row
                       if c.value == 'Radio q'][0]
        self.assertTrue(header_cell.font.b)
        self.assertTrue(len(ws._images) >= 3)
