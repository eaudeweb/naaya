""" Tests for sections, paragraphs, comments, exports and bulk invitations """

import csv
import io
import os
from datetime import date, timedelta
from unittest import TestCase
from unittest.mock import Mock, patch

from Acquisition import aq_base

import openpyxl
import transaction

from Products.Naaya.NyFolder import addNyFolder
from Products.Naaya.tests.NaayaFunctionalTestCase import \
    NaayaFunctionalTestCase
from Products.NaayaCore.EmailTool.EmailTool import divert_mail
from naaya.core.zope2util import XLSX_CONTENT_TYPE

from naaya.content.talkback.comment_item import addComment, cleanup_message
from naaya.content.talkback.tbconsultation_item import \
    addNyTalkBackConsultation
from naaya.content.talkback.tests import (anonymous_contributor,
                                          invite_contributor,
                                          use_temp_sites_log,
                                          user_contributor)

DATA_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'data')
CONS_URL = 'http://localhost/portal/myfolder/test-consultation'


class SanitizeTestCase(TestCase):

    def test_script_and_style_removed_with_content(self):
        self.assertEqual(
            cleanup_message('a<script>alert(1)</script>b'
                            '<style>p {color: red}</style>c'), 'abc')
        self.assertEqual(cleanup_message('<SCRIPT>x</SCRIPT>ok'), 'ok')

    def test_disallowed_tags_and_attributes(self):
        self.assertEqual(
            cleanup_message('<p onclick="x()">t <iframe src="x"></iframe>'
                            '<b>bold</b></p>'),
            '<p>t <b>bold</b></p>')
        self.assertEqual(cleanup_message('<!-- hidden -->text'), 'text')

    def test_links(self):
        self.assertEqual(
            cleanup_message('<a href="http://example.com/" '
                            'onmouseover="x()">x</a>'),
            '<a href="http://example.com/" rel="nofollow" '
            'class="external">x</a>')
        self.assertEqual(
            cleanup_message('<a href="javascript:alert(1)">x</a>'),
            '<a rel="nofollow" class="external">x</a>')
        self.assertEqual(
            cleanup_message('see http://example.com/page'),
            'see <a href="http://example.com/page" rel="nofollow" '
            'class="external">http://example.com/page</a>')

    def test_alignment_and_colours_kept(self):
        self.assertEqual(
            cleanup_message('<p align="center" style="text-align: right; '
                            'position: fixed">a <span style="color: #f00; '
                            'background-color: yellow; font-size: 90px">b'
                            '</span></p>'),
            '<p align="center" style="text-align: right;">a <span '
            'style="color: #f00; background-color: yellow;">b</span></p>')
        self.assertEqual(
            cleanup_message('<p style="background: url(javascript:x)">x</p>'),
            '<p style="">x</p>')

    def test_trim_empty_paragraphs(self):
        self.assertEqual(
            cleanup_message('<p>&nbsp;</p><p>text &amp; more</p><p> </p>'),
            '<p>text &amp; more</p>')


class TalkbackTestCase(NaayaFunctionalTestCase):

    def afterSetUp(self):
        use_temp_sites_log(self)
        addNyFolder(self.portal, 'myfolder', contributor='admin',
                    submitted=1)
        start_date = (date.today() - timedelta(days=1)).strftime('%d/%m/%Y')
        end_date = (date.today() + timedelta(days=10)).strftime('%d/%m/%Y')
        addNyTalkBackConsultation(self.portal.myfolder,
                                  title="Test consultation",
                                  start_date=start_date, end_date=end_date,
                                  contributor='admin', submitted=1)
        self.consultation = self.portal.myfolder['test-consultation']
        self.consultation.addSection(
            id='sec', title='Section',
            body='<p>First</p><p>Second</p><p>Third</p>')
        transaction.commit()
        self.section = self.consultation['sec']
        self.diverted_mail = divert_mail()

    def beforeTearDown(self):
        divert_mail(False)
        self.portal.manage_delObjects(['myfolder'])
        transaction.commit()

    def paragraph_bodies(self):
        return [p.body for p in self.section.get_paragraphs()]


class SectionParagraphTestCase(TalkbackTestCase):

    def test_add_section(self):
        self.assertEqual(self.paragraph_bodies(),
                         ['<p>First</p>', '<p>Second</p>', '<p>Third</p>'])
        self.consultation.addSection(title='Unsplit', skip_splitting='on',
                                     body='<p>One</p><p>Two</p>')
        self.consultation.addSection(title='Sec', body='<p>x</p>')
        # sections without a saved order are sorted by title
        self.assertEqual(
            [s.getId() for s in self.consultation.list_sections()],
            ['sec-1', 'sec', 'unsplit'])
        self.assertEqual(
            [p.body for p in self.consultation['unsplit'].get_paragraphs()],
            ['<p>One</p><p>Two</p>'])

    def test_sort_and_delete_sections(self):
        self.consultation.addSection(id='other', title='A', body='<p>x</p>')
        self.consultation.save_sort_order(['other', 'sec'])
        self.assertEqual(
            [s.getId() for s in self.consultation.list_sections()],
            ['other', 'sec'])
        self.consultation.delete_sections([])
        self.assertEqual(len(self.consultation.list_sections()), 2)
        self.consultation.delete_sections(['other'])
        self.assertEqual(
            [s.getId() for s in self.consultation.list_sections()], ['sec'])

    def test_split_merge_move_remove(self):
        self.login()
        self.section['000'].split_body(body_0='<p>Fir</p>',
                                       body_1='<p>st</p>')
        self.assertEqual(self.paragraph_bodies(),
                         ['<p>Fir</p>', '<p>st</p>', '<p>Second</p>',
                          '<p>Third</p>'])
        new_id = self.section.paragraph_ids[1]

        comment_id = addComment(self.section[new_id],
                                contributor=user_contributor('contributor'),
                                message='on the split paragraph')
        request = self.portal.REQUEST
        self.section['000'].merge_down(request)
        self.assertEqual(self.paragraph_bodies(),
                         ['<p>Fir</p><p>st</p>', '<p>Second</p>',
                          '<p>Third</p>'])
        self.assertEqual(self.section['000'].objectIds(), [comment_id])
        self.assertFalse(new_id in self.section.objectIds())

        self.section['000'].move_down(request)
        self.assertEqual(self.paragraph_bodies(),
                         ['<p>Second</p>', '<p>Fir</p><p>st</p>',
                          '<p>Third</p>'])

        self.section['002'].remove_paragraph(request)
        self.assertEqual(self.paragraph_bodies(),
                         ['<p>Second</p>', '<p>Fir</p><p>st</p>'])
        self.assertEqual(self.section.comment_count(), 1)
        self.logout()

    def test_ensure_paragraph_ids(self):
        self.section.paragraph_ids.append('missing-1')
        self.section.paragraph_ids.append('missing-2')
        self.section._ensure_paragraph_ids()
        self.assertEqual(self.section.paragraph_ids, ['000', '001', '002'])

    def test_pages_render(self):
        addComment(self.section['000'],
                   contributor=user_contributor('contributor'),
                   message='a comment')
        self.consultation.addSection(id='next', title='Next',
                                     body='<p>n</p>')
        transaction.commit()
        self.browser_do_login('admin', '')
        for url, text in [
                (CONS_URL, 'Test consultation'),
                (CONS_URL + '/edit_html', 'Test consultation'),
                (CONS_URL + '/section_add_html', 'frmAdd'),
                (CONS_URL + '/sec', 'Second'),
                (CONS_URL + '/sec?inline_comments=on', 'a comment'),
                (CONS_URL + '/sec/edit_html', 'Second'),
                (CONS_URL + '/sec/000', 'a comment'),
                (CONS_URL + '/sec/000/embedded_html', 'a comment'),
                (CONS_URL + '/sec/000/edit_html', 'First'),
                (CONS_URL + '/sec/000/split_body', 'First'),
                (CONS_URL + '/admin_comments', 'Contributor Test'),
                (CONS_URL + '/admin_comments/comments_table_html',
                 'a comment'),
                (CONS_URL + '/invitations/create', 'invite_bulk'),
                (CONS_URL + '/invitations/send_email', 'Send'),
                (CONS_URL + '/invitations/saved_emails', 'Test consultation'),
                ]:
            self.browser.go(url)
            self.assertEqual(self.browser.get_code(), 200, url)
            html = self.browser.get_html()
            self.assertFalse('Error type' in html, url)
            self.assertTrue(text in html, url)
        self.browser_do_logout()

    def test_delete_sections_page(self):
        self.browser_do_login('admin', '')
        self.browser.go(CONS_URL + '/delete_sections')
        self.assertEqual(self.browser.get_code(), 200)
        self.assertEqual(len(self.consultation.list_sections()), 1)
        self.browser_do_logout()


class CommentsTestCase(TalkbackTestCase):

    def test_delete_comment(self):
        comment_id = addComment(self.section['000'],
                                contributor=user_contributor('contributor'),
                                message='to be deleted')
        transaction.commit()
        self.browser_do_login('admin', '')
        url = CONS_URL + '/sec/000/delete_comment?comment_id=' + comment_id
        self.browser.go(url)
        self.assertTrue('Are you sure you want to delete'
                        in self.browser.get_html())
        self.assertEqual(self.section['000'].objectIds(), [comment_id])
        form = [f for f in self.browser._browser.forms()
                if 'comment_id' in f._form.fields][0]
        self.browser.clicked(form, form.find_control('comment_id'))
        self.browser.submit()
        self.assertEqual(self.section['000'].objectIds(), [])
        self.browser_do_logout()

    def test_cannot_delete_comment_with_replies(self):
        paragraph = self.section['000']
        parent_id = addComment(paragraph,
                               contributor=user_contributor('contributor'),
                               message='parent')
        addComment(paragraph, contributor=user_contributor('contributor'),
                   message='reply', reply_to=parent_id)
        self.login()
        self.assertFalse(paragraph[parent_id].check_del_permissions())
        self.assertEqual(len(paragraph.get_comment_tree()), 1)
        self.assertEqual(
            len(paragraph.get_comment_tree()[0]['children']), 1)
        self.logout()

    def test_anonymous_and_invited_contributor_info(self):
        paragraph = self.section['000']
        anon_id = addComment(paragraph,
                             contributor=anonymous_contributor(
                                 'Gigi', 'gigi@example.com', 'Gigi Inc'),
                             message='anonymous')
        self.consultation.invitations._send_invitation(
            name='The Invitee', email='invitee@example.com',
            organization='Org', notes='', message='',
            inviter_userid='contributor', inviter_name='Contributor Test')
        invitation = list(self.consultation.invitations._invites.values())[0]
        inv_id = addComment(paragraph,
                            contributor=invite_contributor(invitation),
                            message='invited')
        anon_info = paragraph[anon_id].get_contributor_info()
        self.assertEqual(anon_info['display_name'],
                         'Gigi - Gigi Inc (not authenticated)')
        self.assertTrue(anon_info['anonymous'])
        inv_info = paragraph[inv_id].get_contributor_info()
        self.assertEqual(inv_info['display_name'],
                         'The Invitee (invited by Contributor Test)')
        self.assertTrue(inv_info['invited'])
        self.assertEqual(inv_info['email'], 'invitee@example.com')

    def test_empty_comment_rejected(self):
        self.browser_do_login('admin', '')
        self.browser.go(CONS_URL + '/sec/000')
        form = self.browser.get_form('frmAdd')
        form['message:utf8:string'] = '<p>&nbsp;</p>'
        self.browser.clicked(form, form.find_control('message:utf8:string'))
        self.browser.submit()
        self.assertEqual(self.section['000'].objectIds(), [])
        self.assertTrue('The comment field cannot be empty.'
                        in self.browser.get_html())
        self.browser_do_logout()


class CommentsExportTestCase(TalkbackTestCase):

    def afterSetUp(self):
        super(CommentsExportTestCase, self).afterSetUp()
        paragraph = self.section['000']
        self.parent_id = addComment(
            paragraph, contributor=user_contributor('contributor'),
            message='<p>First comment</p>')
        addComment(paragraph, contributor=anonymous_contributor('Gigi'),
                   message='<p>Reply</p>', reply_to=self.parent_id)
        transaction.commit()
        self.browser_do_login('admin', '')

    def beforeTearDown(self):
        self.browser_do_logout()
        super(CommentsExportTestCase, self).beforeTearDown()

    def export(self, file_type):
        self.browser.go(CONS_URL + '/admin_comments/export?file_type=' +
                        file_type + '&as_attachment:boolean=True')
        self.assertEqual(self.browser.get_code(), 200)
        return self.browser.result.content

    def test_csv(self):
        data = self.export('CSV')
        self.assertTrue(self.browser_get_header('content-type')
                        .startswith('text/csv'))
        self.assertTrue('comments.csv' in
                        self.browser_get_header('content-disposition'))
        rows = list(csv.reader(io.StringIO(data.decode('utf-8-sig'))))
        self.assertEqual(rows[2][0], 'Section')
        self.assertEqual(rows[3][5], 'First comment')
        self.assertEqual(rows[3][3], '1')
        self.assertEqual(rows[4][4], self.parent_id)

    def test_excel(self):
        data = self.export('Excel')
        self.assertEqual(self.browser_get_header('content-type'),
                         XLSX_CONTENT_TYPE)
        self.assertTrue('comments.xlsx' in
                        self.browser_get_header('content-disposition'))
        ws = openpyxl.load_workbook(io.BytesIO(data)).worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        self.assertEqual(rows[0][0], 'Consultation deadline')
        self.assertEqual(rows[2][:3], ['Section', 'Paragraph', 'Message Id'])
        self.assertEqual(rows[3][0], 'Section')
        self.assertEqual(rows[3][5], 'First comment')
        self.assertEqual(rows[3][6], 'contributor (Contributor Test)')
        self.assertEqual(ws.cell(row=4, column=7).hyperlink.target,
                         'https://www.eionet.europa.eu/users/contributor')
        self.assertEqual(rows[4][5], 'Reply')
        self.assertEqual(rows[4][6], 'Gigi (not authenticated)')

    def test_custom_excel(self):
        data = self.export('CustomExcel')
        self.assertEqual(self.browser_get_header('content-type'),
                         XLSX_CONTENT_TYPE)
        ws = openpyxl.load_workbook(io.BytesIO(data)).worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        self.assertEqual(rows[0], ['Section', 'Message', 'Contributor',
                                   'Date'])
        self.assertEqual(rows[1][:3], ['Section', 'First comment',
                                       'contributor (Contributor Test)'])
        self.assertEqual(rows[2][0], 'EEA Comments')

    def test_atom_feed(self):
        self.browser.go(CONS_URL + '/comments_atom')
        self.assertTrue('First comment' in self.browser.result.text)


class BulkInvitationsTestCase(TalkbackTestCase):

    def upload(self, filename, button):
        self.browser.go(CONS_URL + '/invitations/create')
        form = self.browser.get_form('invite_bulk')
        with open(os.path.join(DATA_DIR, filename), 'rb') as f:
            data = f.read()
        form.find_control('input_file').add_file(
            io.BytesIO(data), filename=filename,
            content_type='application/octet-stream')
        button_el = form._form.xpath(
            './/input[@name="do"][@value="%s"]' % button)[0]
        self.browser.clicked(form, button_el)
        self.browser.submit()

    def check_import(self, filename):
        self.browser_do_login('admin', '')
        self.upload(filename, 'Preview')
        self.assertEqual(len(self.consultation.invitations._invites), 0)
        self.assertTrue('Hello Ana' in self.browser.get_html())

        self.upload(filename, 'Send invitation')
        invites = sorted(self.consultation.invitations._invites.values(),
                         key=lambda i: i.name)
        self.assertEqual(
            [(i.name, i.email, i.organization, i.notes, i.inviter_userid)
             for i in invites],
            [('Ana Pop', 'ana@example.com', 'Org A', 'Note A', 'admin'),
             ('Ion Ionescu', 'ion@example.com', 'Org B', '', 'admin')])
        self.assertEqual(sorted(m[1][0] for m in self.diverted_mail),
                         ['ana@example.com', 'ion@example.com'])
        self.browser_do_logout()

    def test_import_xlsx(self):
        self.check_import('invitations.xlsx')

    def test_import_xls(self):
        self.check_import('invitations.xls')

    def test_import_invalid_file(self):
        self.browser_do_login('admin', '')
        self.browser.go(CONS_URL + '/invitations/create')
        form = self.browser.get_form('invite_bulk')
        form.find_control('input_file').add_file(
            io.BytesIO(b'not a spreadsheet'), filename='x.xlsx',
            content_type='application/octet-stream')
        self.browser.clicked(form, form._form.xpath(
            './/input[@name="do"][@value="Preview"]')[0])
        self.browser.submit()
        self.assertTrue('Error reading the spreadsheet.'
                        in self.browser.get_html())
        self.browser_do_logout()

    def test_import_requires_permission(self):
        self.browser.go(CONS_URL + '/invitations/bulk_create')
        self.assertAccessDenied()

    def test_export_invitations(self):
        create_invitation = self.consultation.invitations._create_invitation
        key = create_invitation(inviter_userid='contributor', name='Ana',
                                email='ana@example.com', organization='Org',
                                notes='')
        transaction.commit()
        self.browser_do_login('admin', '')
        self.browser.go(CONS_URL + '/invitations/invitations_export'
                        '?headers=Name,Invited%20by,Private%20URL'
                        '&keys=name,name_from_userid,private_url')
        self.assertEqual(self.browser_get_header('content-type'),
                         XLSX_CONTENT_TYPE)
        self.assertTrue('consultation_invitations.xlsx' in
                        self.browser_get_header('content-disposition'))
        wb = openpyxl.load_workbook(io.BytesIO(self.browser.result.content))
        rows = [list(r) for r in
                wb['Active Invitations'].iter_rows(values_only=True)]
        self.assertEqual(rows, [
            ['Name', 'Invited by', 'Private URL'],
            ['Ana', 'Contributor Test',
             CONS_URL + '/invitations/welcome?key=' + key]])
        self.browser_do_logout()

    def test_saved_emails(self):
        self.consultation.invitations._send_invitation(
            name='Ana', email='ana@example.com', organization='Org',
            notes='', message='Please review', inviter_userid='admin',
            inviter_name='Admin')
        transaction.commit()
        self.browser_do_login('admin', '')
        self.browser.go(CONS_URL + '/invitations/saved_emails')
        self.assertTrue('ana@example.com' in self.browser.get_html())
        self.browser.go(CONS_URL + '/invitations/saved_emails_export'
                        '?headers=Subject,Recipients&keys=subject,recipients')
        self.assertEqual(self.browser_get_header('content-type'),
                         XLSX_CONTENT_TYPE)
        self.assertTrue('consultation_invitation_emails.xlsx' in
                        self.browser_get_header('content-disposition'))
        wb = openpyxl.load_workbook(io.BytesIO(self.browser.result.content))
        rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        self.assertEqual(rows[0], ['Subject', 'Recipients'])
        self.assertEqual(len(rows), 2)
        self.assertTrue('ana@example.com' in rows[1][1])
        self.browser_do_logout()


class OnBehalfTestCase(TalkbackTestCase):

    def bulk_preview(self, roles):
        container = self.consultation.invitations
        cls = aq_base(container).__class__
        request = Mock()
        request.REQUEST_METHOD = 'POST'
        request.AUTHENTICATED_USER.getRoles.return_value = roles
        with open(os.path.join(DATA_DIR, 'invitations.xlsx'), 'rb') as f:
            request.form = {'do': 'Preview', 'on_behalf': 'contributor',
                            'input_file': io.BytesIO(f.read())}
        with patch.object(cls, '_create_html', Mock(return_value='page')) \
                as create_html, \
                patch.object(cls, '_get_inviter_info',
                             Mock(return_value=('site_admin', 'SA'))), \
                patch.object(aq_base(self.portal).__class__,
                             'setSessionErrorsTrans', Mock()) as errors:
            container.bulk_create(request)
        return create_html, errors

    def test_on_behalf_rejected_for_non_managers(self):
        create_html, errors = self.bulk_preview(['Administrator',
                                                 'Authenticated'])
        errors.assert_called_once_with(
            'Only managers can invite on behalf of another user')
        self.assertEqual(create_html.call_args[1]['previews'], [])

    def test_on_behalf_allowed_for_managers(self):
        create_html, errors = self.bulk_preview(['Manager', 'Authenticated'])
        self.assertFalse(errors.called)
        previews = create_html.call_args[1]['previews']
        self.assertEqual(len(previews), 2)
        self.assertTrue('invited by Contributor Test'
                        in previews[0]['preview_attribution'])


class CommentSessionTestCase(TalkbackTestCase):

    def test_contributor_cleared_after_comment(self):
        paragraph = self.section['000']
        request = Mock()
        request.form = {'message': '<p>Hello</p>',
                        'contributor': {'name': 'Gigi'}}
        request.get.side_effect = lambda key, default=None: default
        self.login()
        with patch.object(aq_base(self.portal).__class__, 'delSession',
                          Mock()) as delete:
            paragraph.addComment(request)
        self.logout()
        deleted = [c[0][0] for c in delete.call_args_list]
        self.assertIn('contributor', deleted)
        self.assertNotIn('username', deleted)
